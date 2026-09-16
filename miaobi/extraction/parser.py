from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
from typing import Iterable

import chardet
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from miaobi.extraction.models import ParsedDocument, ParsedElement


TEXT_SUFFIXES = {".txt", ".md", ".rst", ".yaml", ".yml", ".xml", ".json", ".jsonl"}


def _stable_id(filename: str, path: str, text: str) -> str:
    return hashlib.sha256(f"{filename}:{path}:{text}".encode("utf-8")).hexdigest()[:24]


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        detected = chardet.detect(data)
        encoding = str(detected.get("encoding") or "").strip()
        confidence = float(detected.get("confidence") or 0)
        if not encoding or confidence < 0.55:
            raise ValueError("文件字符编码无法可靠识别，请转换为 UTF-8 后重试")
        return data.decode(encoding)


def _element(filename: str, kind: str, text: str, path: str, page: int | None = None) -> ParsedElement:
    clean = text.strip()
    return ParsedElement(
        element_id=_stable_id(filename, path, clean),
        element_type=kind,
        text=clean,
        structural_path=path,
        page_number=page,
    )


def _nonempty(items: Iterable[ParsedElement]) -> tuple[ParsedElement, ...]:
    return tuple(item for item in items if item.text.strip())


def parse_document_bytes(filename: str, content_type: str | None, data: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    parser = ""
    elements: tuple[ParsedElement, ...]

    if suffix == ".pdf":
        parser = "pypdf"
        reader = PdfReader(io.BytesIO(data))
        elements = _nonempty(
            _element(filename, "page", page.extract_text() or "", f"pages/{index}", index + 1)
            for index, page in enumerate(reader.pages)
        )
    elif suffix == ".docx":
        parser = "python-docx"
        document = DocxDocument(io.BytesIO(data))
        rows: list[ParsedElement] = []
        for index, paragraph in enumerate(document.paragraphs):
            if paragraph.text.strip():
                rows.append(_element(filename, "paragraph", paragraph.text, f"paragraphs/{index}"))
        for table_index, table in enumerate(document.tables):
            table_text = "\n".join("\t".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            rows.append(_element(filename, "table", table_text, f"tables/{table_index}"))
        elements = _nonempty(rows)
    elif suffix == ".xlsx":
        parser = "openpyxl"
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                text = "\t".join("" if value is None else str(value) for value in values)
                rows.append(_element(filename, "record", text, f"sheets/{sheet.title}/rows/{row_index}"))
        elements = _nonempty(rows)
    elif suffix == ".csv":
        parser = "csv"
        text = _decode(data)
        rows = [
            _element(filename, "record", "\t".join(row), f"rows/{index}")
            for index, row in enumerate(csv.reader(io.StringIO(text)), start=1)
        ]
        elements = _nonempty(rows)
    elif suffix in {".html", ".htm"}:
        parser = "beautifulsoup"
        soup = BeautifulSoup(_decode(data), "html.parser")
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        elements = _nonempty([_element(filename, "text", soup.get_text("\n", strip=True), "document")])
    elif suffix in TEXT_SUFFIXES:
        parser = "text"
        text = _decode(data)
        if suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError as exc:
                raise ValueError(f"JSON 格式不正确：第 {exc.lineno} 行") from exc
        elements = _nonempty([_element(filename, "text", text, "document")])
    else:
        raise ValueError(f"MVP 暂不支持 {suffix or '无扩展名'} 文件")

    if not elements:
        raise ValueError("文件没有解析出可用文字；扫描 PDF 后续接入 OCR")
    return ParsedDocument(
        filename=Path(filename).name,
        content_type=content_type or "application/octet-stream",
        sha256=hashlib.sha256(data).hexdigest(),
        parser=parser,
        elements=elements,
    )

